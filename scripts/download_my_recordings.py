#!/usr/bin/env python3
"""One-click Tencent Meeting "My Recordings" downloader via CDP proxy.

Usage:
    python3 download_my_recordings.py                          # download all
    python3 download_my_recordings.py --start-from 2026-03-01  # from March 2026
    python3 download_my_recordings.py --list-only              # enumerate only
    python3 download_my_recordings.py --resume                 # skip existing files

Prerequisites (Linux):
    - scripts/start_chrome.sh 已启动 headless Chrome（默认 127.0.0.1:9333）
    - 该 Chrome profile 已扫码登录 meeting.tencent.com（scripts/login_via_qr.py tencent）
    - （旧 macOS 代理模式仍可用：--cdp-mode proxy --cdp-port 3456）
    - openpyxl installed (pip install openpyxl)
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone, timedelta
from http.client import HTTPConnection
from pathlib import Path
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from urllib.error import HTTPError

from recording_naming import (
    RESOURCE_LABELS,
    candidate_media_paths,
    find_existing_complete_path,
    media_filename,
    normalize_resource_type,
    segment_indexes,
)

import linux_env
import cdp_client

CDP_HOST = "localhost"
CDP_PORT = linux_env.CDP_PORT
# native: 直连 Chrome --remote-debugging-port（Linux 默认）
# proxy : 旧 macOS web-access 代理（localhost:3456 的 /new /eval 接口）
CDP_MODE = "native"
OUTPUT_DIR = linux_env.DOWNLOAD_ROOT
API_BASE = "/wemeet-tapi/v2"
COMMON_Q = "c_os_model=web&c_os=web&c_instance_id=5&c_account_corp_id=656648463&c_lang=zh-CN"
REFERER = "https://meeting.tencent.com/user-center/meeting-record"
UA = linux_env.FALLBACK_UA  # 运行时用浏览器真实 UA 覆盖（refresh_auth_from_browser）
# Signed COS URLs (ylz.cos.meeting.tencent.com) reject requests without the
# logged-in session cookie (403). Callers set this to the browser's
# document.cookie so download_file can authenticate the file fetch.
COOKIE = ""

CST = timezone(timedelta(hours=float(os.environ.get("LINK_DL_TZ_OFFSET", "8"))))


# ── CDP proxy layer ──────────────────────────────────────────────────

def cdp_call(method, path, body=None, timeout=30):
    if CDP_MODE == "native":
        return cdp_client.legacy_call(method, path, body, port=CDP_PORT, timeout=timeout)
    conn = HTTPConnection(CDP_HOST, CDP_PORT, timeout=timeout)
    headers = {}
    if body is not None:
        body = body.encode("utf-8") if isinstance(body, str) else body
        headers["Content-Type"] = "application/x-www-form-urlencoded"
    conn.request(method, path, body, headers)
    resp = conn.getresponse()
    raw = resp.read().decode()
    conn.close()
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return raw
    return data


def cdp_eval(target_id, js, timeout=60):
    r = cdp_call("POST", f"/eval?target={target_id}", js, timeout=timeout)
    return r.get("value")


def refresh_auth_from_browser(target_id, cookie_urls=None):
    """COS 签名 URL 需要腾讯会议登录 cookie（含 HttpOnly）+ 浏览器 UA，否则 403。"""
    global COOKIE, UA
    cookie_urls = cookie_urls or ["https://meeting.tencent.com/", "https://ylz.cos.meeting.tencent.com/"]
    if CDP_MODE == "native":
        s = cdp_client.default_browser(CDP_PORT).session(target_id)
        COOKIE = s.cookie_header(cookie_urls)
        UA = s.user_agent()
    else:
        COOKIE = cdp_eval(target_id, "document.cookie") or ""
        UA = linux_env.clean_ua(cdp_eval(target_id, "navigator.userAgent"))
    return COOKIE


def find_or_create_tab():
    targets = cdp_call("GET", "/targets")
    if isinstance(targets, list):
        pages = [t for t in targets if isinstance(t, dict) and t.get("type") == "page" and "meeting.tencent.com/user-center/meeting-record" in t.get("url", "")]
        if pages:
            return pages[0]["targetId"]
    r = cdp_call("GET", "/new?url=https://meeting.tencent.com/user-center/meeting-record")
    if isinstance(r, dict) and "targetId" in r:
        time.sleep(5)
        return r["targetId"]
    time.sleep(5)
    targets = cdp_call("GET", "/targets")
    if isinstance(targets, list):
        for t in targets:
            if isinstance(t, dict) and "meeting.tencent.com/user-center/meeting-record" in t.get("url", ""):
                return t["targetId"]
    raise RuntimeError("Failed to create or find tab")


def ensure_logged_in(target_id):
    url = cdp_eval(target_id, "document.location.href")
    if not url or "login" in str(url):
        print("ERROR: 未登录。先运行 python3 scripts/login_via_qr.py tencent 扫码登录，再重跑。")
        sys.exit(1)
    cookie_check = cdp_eval(target_id, 'document.cookie.indexOf("token_expire_time") > -1 ? "yes" : "no"')
    if cookie_check != "yes" and CDP_MODE == "native":
        names = {c["name"] for c in cdp_client.default_browser(CDP_PORT).session(target_id).cookies(["https://meeting.tencent.com/"])}
        cookie_check = "yes" if "token_expire_time" in names else "no"
    if cookie_check != "yes":
        print("ERROR: 未登录。先运行 python3 scripts/login_via_qr.py tencent 扫码登录，再重跑。")
        sys.exit(1)


# ── Browser-side XHR ─────────────────────────────────────────────────

def browser_xhr(target_id, method, path, body=None, retries=3):
    for attempt in range(retries):
        nonce = f"{''.join(__import__('random').choice(__import__('string').ascii_letters + __import__('string').digits) for _ in range(8))}"
        ts = int(time.time() * 1000)
        qs = f"c_os_model=web&c_os=web&c_timestamp={ts}&c_nonce={nonce}&c_instance_id=5&c_account_corp_id=656648463&rnds={nonce}&c_lang=zh-CN"
        url = f"{path}?{qs}" if "?" not in path else f"{path}&{qs}"
        body_js = json.dumps(json.dumps(body)) if body else "null"
        js = (
            "(function(){"
            "  try {"
            "    var xhr = new XMLHttpRequest();"
            f"    xhr.open('{method}', '{url}', false);"
            "    xhr.setRequestHeader('Content-Type', 'application/json');"
            f"    xhr.send({body_js});"
            "    return JSON.stringify({status: xhr.status, body: xhr.responseText});"
            "  } catch(e) {"
            "    return JSON.stringify({error: e.message});"
            "  }"
            "})()"
        )
        raw = cdp_eval(target_id, js, timeout=30)
        if raw is None:
            if attempt < retries - 1:
                time.sleep(1)
                continue
            raise RuntimeError("CDP eval returned None")
        parsed = json.loads(raw)
        if "error" in parsed:
            if attempt < retries - 1:
                time.sleep(1)
                continue
            raise RuntimeError(f"XHR error: {parsed['error']}")
        data = json.loads(parsed["body"])
        if data.get("code") != 0:
            if attempt < retries - 1:
                time.sleep(2)
                continue
            raise RuntimeError(f"API error code={data.get('code')}: {data.get('msg', '')[:100]}")
        return data


# ── API functions ────────────────────────────────────────────────────

def fetch_record_list(target_id, page_index=1, page_size=30):
    body = {"page_index": page_index, "page_size": page_size, "record_type": "cloud_record"}
    data = browser_xhr(target_id, "POST", f"{API_BASE}/meetlog/dashboard/my-record-list", body)
    records = data.get("data", {}).get("records", [])
    total = int(data.get("data", {}).get("page_info", {}).get("count", 0))
    return records, total


def fetch_record_detail(target_id, encode_record_id):
    path = f"{API_BASE}/meetlog/public/record-detail/get-multi-record-info"
    path += f"?pwd=&auth_share_id={encode_record_id}&uni_record_share_id={encode_record_id}&activity_uid="
    data = browser_xhr(target_id, "GET", path)
    return data.get("data", {}).get("base_infos", [])


def fetch_record_files(target_id, encode_record_id, recording_id, stream_type):
    path = f"{API_BASE}/meetlog/public/record-detail/get-multi-record-file"
    path += f"?record_id={recording_id}&auth_share_id={encode_record_id}&stream_type={stream_type}&pwd=&activity_uid="
    data = browser_xhr(target_id, "GET", path)
    return data.get("data", {}).get("files", [])


def sign_download_url(target_id, encode_record_id, recording_id, resource_id, resource_type):
    path = f"{API_BASE}/meetlog/public/record-detail/sign-multi-record-file"
    path += f"?record_id={recording_id}&resource_id={resource_id}&resource_type={resource_type}&auth_share_id={encode_record_id}&pwd=&activity_uid="
    data = browser_xhr(target_id, "GET", path)
    url = data.get("data", {}).get("url")
    if not url:
        raise RuntimeError("sign returned no URL")
    return url


def fetch_transcription(target_id, encode_record_id, meeting_id, recording_id):
    """逐字稿分页拉全（旧实现只拿第一页 20 段）。"""
    try:
        import tencent_minutes

        pages = tencent_minutes.fetch_all_minutes_pages(
            lambda p: browser_xhr(target_id, "GET", p),
            encode_record_id, meeting_id, recording_id,
        )
        merged = tencent_minutes.merge_minutes_pages(pages)
        if merged and merged.get("paragraphs"):
            return merged
    except Exception as exc:
        print(f"  分页拉逐字稿失败，退回单页: {exc}")
    return _fetch_transcription_first_page(target_id, encode_record_id, meeting_id, recording_id)


def _fetch_transcription_first_page(target_id, encode_record_id, meeting_id, recording_id):
    path = "/wemeet-cloudrecording-webapi/v1/minutes/detail"
    qs = f"c_os_model=web&c_os=web&c_timestamp={int(time.time()*1000)}&c_nonce={'t'*8}&c_instance_id=5&c_account_corp_id=656648463&c_lang=zh-CN"
    qs += f"&mock=1&platform=Web&id={encode_record_id}&meeting_id={meeting_id}&recording_id={recording_id}"
    qs += "&start_pid=0&limit=20&fview=1&minutes_version=0&return_ori=0&return_ori_minutes_translating=1&lang=zh&page_source=record"
    data = browser_xhr(target_id, "GET", f"{path}?{qs}")
    return data.get("minutes")


def format_transcription(minutes):
    if not minutes or "paragraphs" not in minutes:
        return None
    lines = []
    for para in minutes.get("paragraphs", []):
        speaker = para.get("speaker", {})
        name = speaker.get("user_name", "未知")
        start_ms = int(para.get("start_time", 0))
        m, s = divmod(start_ms // 1000, 60)
        h, m = divmod(m, 60)
        ts = f"{h:02d}:{m:02d}:{s:02d}"
        text = ""
        for sent in para.get("sentences", []):
            for word in sent.get("words", []):
                text += word.get("text", "")
        if text.strip():
            lines.append(f"[{ts}] {name}：{text.strip()}")
    return "\n".join(lines) if lines else None


def save_transcription(text, dest):
    dest.write_text(text, encoding="utf-8")
    return "done"


# ── Enumerate ────────────────────────────────────────────────────────

def enumerate_all(target_id, start_from=None, start_to=None):
    print("Enumerating recordings...")
    all_records = []
    seen_eids = set()
    start_ms = int(datetime.strptime(start_from, "%Y-%m-%d").replace(tzinfo=CST).timestamp() * 1000) if start_from else 0
    end_ms = int(datetime.strptime(start_to, "%Y-%m-%d").replace(tzinfo=CST).timestamp() * 1000) if start_to else float("inf")

    page = 0
    while True:
        page += 1
        records, total = fetch_record_list(target_id, page)
        if not records:
            break
        stop = False
        for r in records:
            eid = r.get("encode_record_id", "")
            if eid in seen_eids:
                continue
            st = int(r.get("start_time", 0))
            if start_ms and st < start_ms:
                stop = True
                break
            if start_to and st > end_ms:
                continue
            seen_eids.add(eid)
            all_records.append(r)
        print(f"  Page {page}: {len(records)} records, total={total}, collected={len(all_records)}")
        if stop or len(all_records) >= total:
            break
        time.sleep(0.3)

    print(f"  Found {len(all_records)} recordings across {page} pages")
    return all_records


# ── Download ─────────────────────────────────────────────────────────

def safe_name(value, fallback="untitled"):
    value = str(value or "").strip() or fallback
    value = re.sub(r'[\\/:*?"<>|\r\n\t]+', "_", value)
    value = re.sub(r"\s+", " ", value).strip(" .")
    return value[:120] or fallback


def ts_prefix(milliseconds):
    # 服务器时区常是 UTC；目录名统一按会议所在时区（默认北京时间，LINK_DL_TZ_OFFSET 可改）
    try:
        return datetime.fromtimestamp(int(milliseconds) / 1000, tz=CST).strftime("%Y%m%d_%H%M")
    except Exception:
        return "unknown_time"


def meeting_dir_name(record):
    """Build per-meeting subdirectory name: {YYYYMMDD}_{HHMM}_{title}"""
    ts = ts_prefix(record.get("start_time"))
    title = safe_name(record.get("title", "未命名会议"))
    return f"{ts}_{title}"


def make_filename(record, resource_type, segment_index=None):
    return media_filename(record.get("title"), resource_type, segment_index=segment_index)


def download_file(url, dest, expected_size=0, retries=3):
    if dest.exists() and dest.stat().st_size > 0:
        actual = dest.stat().st_size
        if not expected_size or actual == int(expected_size):
            return "skip"
    tmp = dest.with_suffix(dest.suffix + ".part")
    headers = {"User-Agent": UA, "Referer": REFERER, "Accept": "*/*", "Accept-Encoding": "identity"}
    if COOKIE:
        headers["Cookie"] = COOKIE
    for attempt in range(retries):
        # 腾讯 COS 播放 URL 不带 Range 会 403：新下载也必须 bytes=0-，续传用 bytes=<已下>-
        existing = tmp.stat().st_size if tmp.exists() else 0
        headers["Range"] = f"bytes={existing}-"
        try:
            req = Request(url, headers=headers)
            with urlopen(req, timeout=120) as resp:
                if existing and resp.status == 206:
                    cr = resp.headers.get("Content-Range", "")
                    if not cr.startswith(f"bytes {existing}-"):
                        raise RuntimeError(f"invalid Content-Range: {cr}")
                if existing and resp.status != 206:
                    raise RuntimeError(f"resume request returned HTTP {resp.status}")
                mode = "ab" if existing and resp.status == 206 else "wb"
                if mode == "wb":
                    existing = 0
                with tmp.open(mode) as fh:
                    while True:
                        chunk = resp.read(1024 * 1024)
                        if not chunk:
                            break
                        fh.write(chunk)
            actual = tmp.stat().st_size
            if expected_size and actual != int(expected_size):
                return "incomplete"
            tmp.replace(dest)
            return "done"
        except HTTPError as e:
            if e.code == 416 and existing:
                if expected_size and existing == int(expected_size):
                    tmp.replace(dest)
                    return "done"
                return "error: HTTP 416 before exact expected size"
            if attempt < retries - 1:
                time.sleep(2)
            else:
                return f"error: HTTP {e.code}"
        except Exception as e:
            if attempt < retries - 1:
                time.sleep(2)
            else:
                return f"error: {e}"
    return "error: max retries"


# ── Sign + download batch ────────────────────────────────────────────

def sign_and_download(target_id, records, output_dir, resource_types=None, max_retries=1, download_transcriptions=True):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    total = len(records)

    for idx, rec in enumerate(records):
        eid = rec.get("encode_record_id")
        title = rec.get("title", "untitled")
        print(f"\n[{idx+1}/{total}] {title}")

        # Per-meeting subdirectory
        dir_name = meeting_dir_name(rec)
        meeting_dir = output_dir / dir_name
        meeting_dir.mkdir(parents=True, exist_ok=True)

        try:
            versions = fetch_record_detail(target_id, eid)
        except Exception as e:
            print(f"  SKIP detail: {e}")
            results.append({"title": title, "dir": dir_name, "status": "detail_error", "error": str(e)})
            continue

        meeting_id = rec.get("meeting_info", {}).get("meeting_id", "")
        meeting_code = rec.get("meeting_info", {}).get("meeting_code", "")
        segment_by_recording = segment_indexes(versions)
        for ver in versions:
            rid = ver.get("recording_id")
            segment_index = segment_by_recording.get(rid)
            st = ver.get("stream_type")

            try:
                files = fetch_record_files(target_id, eid, rid, st)
            except Exception as e:
                print(f"  SKIP files (stream {st}): {e}")
                continue

            for f in files:
                rt = normalize_resource_type(f.get("resource_type"))
                if resource_types and rt not in {normalize_resource_type(item) for item in resource_types}:
                    continue
                if not f.get("downloadable"):
                    continue

                filename = make_filename(rec, rt, segment_index)
                dest = meeting_dir / filename
                sz = f.get("size", 0)
                existing = find_existing_complete_path(
                    meeting_dir,
                    title,
                    rt,
                    expected_size=sz,
                    segment_index=segment_index,
                    tolerance=0.95,
                )

                if existing:
                    print(f"  SKIP (exists): {existing.name}")
                    results.append({
                        "title": title,
                        "dir": dir_name,
                        "filename": existing.name,
                        "local_path": str(existing),
                        "resource_type": rt,
                        "segment_index": segment_index,
                        "status": "skip",
                        "size": sz,
                    })
                    continue

                for retry in range(max_retries + 1):
                    try:
                        signed = sign_download_url(target_id, eid, rid, f.get("resource_id"), rt)
                    except Exception as e:
                        print(f"  SIGN error: {e}")
                        break

                    size_mb = int(f.get("size", 0)) / 1024**2
                    print(f"  Downloading {filename} ({size_mb:.1f}MB)...", end=" ", flush=True)
                    status = download_file(signed, dest, f.get("size"))
                    print(status)

                    if status in ("done", "skip"):
                        results.append({
                            "title": title,
                            "dir": dir_name,
                            "filename": filename,
                            "local_path": str(dest),
                            "resource_type": rt,
                            "segment_index": segment_index,
                            "status": status,
                            "size": f.get("size"),
                        })
                        break
                    elif "error" in status and retry < max_retries:
                        print(f"  Retry signing...")
                        continue
                    else:
                        results.append({"title": title, "dir": dir_name, "filename": filename, "status": status})
                        break

        # Download transcription
        if download_transcriptions and meeting_id:
            txt_filename = "转写.txt"
            txt_dest = meeting_dir / txt_filename
            if not txt_dest.exists():
                try:
                    recording_id = versions[0]["recording_id"] if versions else ""
                    if recording_id:
                        minutes = fetch_transcription(target_id, eid, meeting_id, recording_id)
                        text = format_transcription(minutes)
                        if text:
                            save_transcription(text, txt_dest)
                            print(f"  Transcription saved: {txt_filename}")
                            results.append({"title": title, "dir": dir_name, "filename": txt_filename, "status": "done"})
                        else:
                            print(f"  No transcription available")
                    else:
                        print(f"  SKIP transcription: no recording_id")
                except Exception as e:
                    print(f"  Transcription error: {e}")
            else:
                print(f"  SKIP (exists): {txt_filename}")
                results.append({"title": title, "dir": dir_name, "filename": txt_filename, "status": "skip"})

        # Save per-meeting manifest.json
        manifest = {
            "title": title,
            "meeting_code": meeting_code,
            "meeting_id": meeting_id,
            "encode_record_id": eid,
            "start_time": rec.get("start_time"),
            "start_time_iso": datetime.fromtimestamp(int(rec.get("start_time", 0)) / 1000, tz=CST).isoformat() if rec.get("start_time") else None,
            "platform": "tencent",
            "files": [r for r in results if r.get("dir") == dir_name and r.get("status") in ("done", "skip")],
        }
        manifest_path = meeting_dir / "manifest.json"
        manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

        time.sleep(0.2)

    return results


# ── Excel catalog ────────────────────────────────────────────────────

def export_excel(results, output_dir):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("openpyxl not installed, skipping Excel export")
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "录制清单"

    headers = ["序号", "主题", "目录", "文件名", "状态", "大小(MB)"]
    hdr_font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    thin = Border(left=Side("thin", "D9D9D9"), right=Side("thin", "D9D9D9"),
                  top=Side("thin", "D9D9D9"), bottom=Side("thin", "D9D9D9"))
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    done_fill = PatternFill("solid", fgColor="E2EFDA")
    err_fill = PatternFill("solid", fgColor="FCE4EC")

    for i, r in enumerate(results, 2):
        status = r.get("status", "")
        ws.cell(i, 1, i - 1).border = thin
        ws.cell(i, 2, r.get("title", "")).border = thin
        ws.cell(i, 3, r.get("dir", "")).border = thin
        ws.cell(i, 4, r.get("filename", "")).border = thin
        ws.cell(i, 5, status).border = thin
        ws.cell(i, 6, round(int(r.get("size", 0)) / 1024**2, 1) if r.get("size") else "").border = thin
        fill = done_fill if status == "done" else (err_fill if "error" in status else None)
        if fill:
            for c in range(1, 7):
                ws.cell(i, c).fill = fill

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 45
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 12
    ws.auto_filter.ref = f"A1:F{len(results)+1}"
    ws.freeze_panes = "A2"

    path = output_dir / "录制目录.xlsx"
    wb.save(path)
    print(f"Catalog saved: {path}")


# ── Main ─────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="One-click Tencent Meeting recordings downloader")
    parser.add_argument("--start-from", default="", help="Start date filter (YYYY-MM-DD)")
    parser.add_argument("--start-to", default="", help="End date filter (YYYY-MM-DD)")
    parser.add_argument("--output-dir", default=str(OUTPUT_DIR))
    parser.add_argument("--resource-types", default="", help="Comma-separated: 0,1,2 (default: all)")
    parser.add_argument("--list-only", action="store_true", help="Only enumerate, don't download")
    parser.add_argument("--resume", action="store_true", help="Skip already-downloaded files (default)")
    parser.add_argument("--cdp-port", type=int, default=linux_env.CDP_PORT)
    parser.add_argument("--cdp-mode", choices=["native", "proxy"], default="native")
    parser.add_argument("--max-retries", type=int, default=1, help="Re-sign retries on download failure")
    parser.add_argument("--no-transcription", action="store_true", help="Skip transcription download")
    args = parser.parse_args()

    global CDP_PORT, CDP_MODE
    CDP_PORT = args.cdp_port
    CDP_MODE = args.cdp_mode

    output_dir = Path(args.output_dir)
    rt_filter = set(int(x) for x in args.resource_types.split(",") if x.strip()) if args.resource_types else None

    # Step 1: Connect to CDP
    print("=== Step 1: Connecting to browser ===")
    target_id = find_or_create_tab()
    ensure_logged_in(target_id)
    refresh_auth_from_browser(target_id)
    print("  Connected and logged in.")

    # Step 2: Enumerate
    print("\n=== Step 2: Enumerating recordings ===")
    records = enumerate_all(target_id, args.start_from or None, args.start_to or None)
    if not records:
        print("No recordings found.")
        cdp_call("GET", f"/close?target={target_id}")
        return

    if args.list_only:
        for i, r in enumerate(records, 1):
            dt = datetime.fromtimestamp(int(r["start_time"]) / 1000, tz=CST)
            print(f"  {i}. [{dt:%Y-%m-%d %H:%M}] {r['title']} ({r.get('meeting_info',{}).get('meeting_code','')})")
        print(f"\nTotal: {len(records)} recordings. Use without --list-only to download.")
        cdp_call("GET", f"/close?target={target_id}")
        return

    # Step 3: Sign + download
    print(f"\n=== Step 3: Downloading to {output_dir} ===")
    results = sign_and_download(target_id, records, output_dir, rt_filter, args.max_retries, download_transcriptions=not args.no_transcription)

    # Step 4: Cleanup & export
    print("\n=== Step 4: Finalizing ===")
    cdp_call("GET", f"/close?target={target_id}")

    done = sum(1 for r in results if r["status"] == "done")
    skip = sum(1 for r in results if r["status"] == "skip")
    errors = sum(1 for r in results if "error" in r["status"])

    export_excel(results, output_dir)

    print(f"\n=== Complete ===")
    print(f"  Downloaded: {done}")
    print(f"  Skipped (existed): {skip}")
    print(f"  Errors: {errors}")
    print(f"  Total files: {len(results)}")
    total_size = sum(int(r.get("size", 0)) for r in results if r["status"] in ("done", "skip"))
    print(f"  Total size: {total_size / 1024**3:.1f} GB")


if __name__ == "__main__":
    main()
